import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Establece color de fondo de una celda"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Establece márgenes internos de una celda"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_study_guide():
    doc = Document()

    # Configuración de márgenes (1 pulgada = 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Colores temáticos
    COLOR_PRIMARY = RGBColor(30, 58, 138)     # Indigo oscuro #1e3a8a
    COLOR_SECONDARY = RGBColor(59, 130, 246)  # Azul #3b82f6
    COLOR_ACCENT = RGBColor(16, 185, 129)     # Esmeralda #10b981
    COLOR_TEXT = RGBColor(30, 41, 59)         # Pizarra #1e293b
    COLOR_MUTED = RGBColor(100, 116, 139)     # Gris #64748b

    # =========================================================================
    # PORTADA
    # =========================================================================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(8)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("GUÍA MAESTRA DE ESTUDIO Y EXPOSICIÓN")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(18)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Sistema de Matrícula Universitaria en Arquitectura MVC\nSpring Boot 4 • Java 21 • MySQL 8.0 • REST API")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(15)
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Separador decorativo
    line_p = doc.add_paragraph()
    line_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_line = line_p.add_run("____________________________________________________")
    run_line.font.color.rgb = COLOR_MUTED
    line_p.paragraph_format.space_after = Pt(24)

    # Cuadro informativo de portada
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Proyecto:", "Sistema de Gestión y Matrícula Académica (sistema_matricula)"),
        ("Arquitectura:", "Modelo - Vista - Controlador (MVC) en Capas"),
        ("Base de Datos:", "MySQL 8.0 (12 Tablas Normalizadas en 3FN)"),
        ("Backend / Framework:", "Java 21 LTS + Spring Boot 4 + Spring Data JPA"),
        ("Objetivo:", "Material de preparación, defensa técnica y exposición final")
    ]
    for i, (k, v) in enumerate(info_data):
        row = info_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(10.5)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = COLOR_TEXT

    doc.add_page_break()

    # =========================================================================
    # SECCIÓN 1: INTRODUCCIÓN Y ALCANCE
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. INTRODUCCIÓN Y ALCANCE DEL PROYECTO")
    r.font.color.rgb = COLOR_PRIMARY

    p = doc.add_paragraph()
    p.add_run("El ").font.color.rgb = COLOR_TEXT
    r = p.add_run("Sistema de Matrícula Universitaria")
    r.bold = True
    p.add_run(" es una solución de software empresarial diseñada para automatizar el proceso integral de inscripción curricular de estudiantes en una institución de educación superior. Resuelve problemas comunes como la sobreventa de cupos, la asignación inadecuada de materias y la desorganización de los cobros académicos.").font.color.rgb = COLOR_TEXT

    doc.add_heading(level=2).add_run("1.1 Requisitos Académicos Cumplidos")
    reqs = [
        ("Mínimo 10 Tablas en Base de Datos: ", "Se implementaron 12 tablas relacionales normalizadas en 3FN, superando el mínimo exigido."),
        ("Patrón Arquitectónico MVC Estricto: ", "Separación total de responsabilidades en Modelo (Entity/DAO), Vista (Frontend SPA) y Controlador (REST Controller)."),
        ("Motor de Base de Datos MySQL: ", "Base de datos 'sistema_matricula' configurada con motor InnoDB, claves foráneas e integridad referencial."),
        ("Gestor de Versiones Git & GitHub: ", "Proyecto versionado con historial de commits y listo para despliegue."),
        ("Operaciones Transaccionales: ", "Lógica de negocio robusta con anotaciones @Transactional para asegurar que las vacantes y los pagos se descuenten de forma atómica.")
    ]
    for title, desc in reqs:
        p = doc.add_paragraph(style='List Bullet')
        r_b = p.add_run(title)
        r_b.bold = True
        r_b.font.color.rgb = COLOR_PRIMARY
        r_d = p.add_run(desc)
        r_d.font.color.rgb = COLOR_TEXT

    # =========================================================================
    # SECCIÓN 2: ARQUITECTURA MVC EN SPRING BOOT
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("2. ARQUITECTURA MVC Y PATRONES DE DISEÑO")
    r.font.color.rgb = COLOR_PRIMARY

    p = doc.add_paragraph()
    p.add_run("El sistema se organizó bajo la filosofía de ").font.color.rgb = COLOR_TEXT
    r = p.add_run("Arquitectura en Capas (Layered Architecture)")
    r.bold = True
    p.add_run(", donde cada componente tiene una única razón para cambiar (Principio de Responsabilidad Única - SRP).").font.color.rgb = COLOR_TEXT

    # Tabla explicativa de capas
    capas_table = doc.add_table(rows=6, cols=3)
    capas_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Capa", "Componentes / Anotaciones", "Responsabilidad Principal"]
    col_widths = [Inches(1.5), Inches(2.2), Inches(2.8)]
    
    for j, h in enumerate(headers):
        cell = capas_table.rows[0].cells[j]
        cell.width = col_widths[j]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    data_capas = [
        ("1. Vista (View)", "Frontend SPA\nHTML5, TailwindCSS, Fetch API", "Interfaz interactiva que presenta la información al usuario y envía solicitudes asíncronas JSON al backend."),
        ("2. Controlador (Controller)", "@RestController\n@RequestMapping\n@GetMapping, @PostMapping", "Punto de entrada de las peticiones HTTP. Valida parámetros, delega a los servicios y retorna respuestas ApiResponse<T> con códigos HTTP (200, 201, 400)."),
        ("3. Servicio (Service)", "Interfaces (IMatriculaService)\nImplementaciones (@Service)\n@Transactional", "Contiene el núcleo de la lógica de negocio. Valida cupos, suma créditos, genera el código de matrícula y garantiza la atomicidad de transacciones."),
        ("4. Repositorio (DAO)", "Interfaces JpaRepository<T, ID>\nConsultas JPQL (@Query)", "Capa de persistencia. Ejecuta operaciones CRUD y búsquedas avanzadas en MySQL sin requerir SQL manual gracias a Spring Data JPA."),
        ("5. Modelo (Entity)", "Clases Java POJO\n@Entity, @Table, @Id\n@ManyToOne, @OneToMany", "Representa las entidades del negocio mapeadas exactamente a las tablas relacionales de la base de datos.")
    ]

    for i, (c1, c2, c3) in enumerate(data_capas, start=1):
        row = capas_table.rows[i]
        bg = "F8FAFC" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate([c1, c2, c3]):
            cell = row.cells[j]
            cell.width = col_widths[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if j == 0:
                r.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            else:
                r.font.color.rgb = COLOR_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    doc.add_heading(level=2).add_run("2.1 El Patrón DTO (Data Transfer Object)")
    p = doc.add_paragraph()
    p.add_run("Uno de los puntos clave del diseño fue la implementación de DTOs (como ").font.color.rgb = COLOR_TEXT
    p.add_run("MatriculaRequestDTO").bold = True
    p.add_run(" y ").font.color.rgb = COLOR_TEXT
    p.add_run("MatriculaResponseDTO").bold = True
    p.add_run("). Esto resolvió el clásico error de ").font.color.rgb = COLOR_TEXT
    p.add_run("Recursión Infinita de Jackson").bold = True
    p.add_run(" que ocurre cuando una entidad bidireccional (Matrícula -> Detalle -> Matrícula) se serializa a JSON, garantizando un contrato de API limpio y desacoplado del modelo interno.").font.color.rgb = COLOR_TEXT

    # =========================================================================
    # SECCIÓN 3: DICCIONARIO Y JUSTIFICACIÓN DE LAS 12 TABLAS
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("3. MODELO DE DATOS Y JUSTIFICACIÓN DE LAS 12 TABLAS")
    r.font.color.rgb = COLOR_PRIMARY

    p = doc.add_paragraph()
    p.add_run("La base de datos ").font.color.rgb = COLOR_TEXT
    p.add_run("sistema_matricula").bold = True
    p.add_run(" fue diseñada con 12 tablas estrechamente interconectadas que reflejan la realidad académica de una universidad:").font.color.rgb = COLOR_TEXT

    tablas_info = [
        ("1. carreras", "Almacena los programas universitarios (Ing. Sistemas, Industrial, Administración, etc.). Posee duración en ciclos y facultad."),
        ("2. planes_estudio", "Representa la versión curricular de una carrera. Una carrera puede tener varios planes a través de los años (ej. Plan 2024)."),
        ("3. periodos_academicos", "Representa los semestres lectivos (2026-I, 2026-II). Permite controlar periodos ACTIVOS, PROXIMOS o CERRADOS."),
        ("4. cursos", "Catálogo de asignaturas (Algoritmos, Base de Datos, IA) vinculadas a un plan de estudio con créditos, horas y ciclo."),
        ("5. docentes", "Catedráticos de la institución con DNI, correo institucional, grado académico (Magister, Doctor) y especialidad."),
        ("6. aulas", "Infraestructura física (Aulas teóricas, Laboratorios de cómputo, Auditorios) con pabellón y capacidad máxima de aforo."),
        ("7. estudiantes", "Padrón oficial de alumnos con DNI, carnet universitario (código único), datos de contacto y carrera inscrita."),
        ("8. secciones", "Grupos ofertados para un curso en un periodo lectivo (ej. Curso SIS101 - Periodo 2026-I - Sección A) con docente y cupo de vacantes."),
        ("9. horarios", "Distribución cronológica por sección indicando día de la semana (LUNES, MARTES...), hora de inicio, hora de fin y aula asignada."),
        ("10. matriculas", "Ficha principal de inscripción que enlaza al estudiante con el periodo académico, calculando total de créditos y costo."),
        ("11. detalles_matricula", "Tabla asociativa que lista cada sección en la que quedó inscrito el alumno, guardando su costo histórico y nota final."),
        ("12. pagos", "Registro financiero de recaudación que vincula el comprobante de pago, número de operación y método (Efectivo, Tarjeta, Yape).")
    ]

    for nom, desc in tablas_info:
        p = doc.add_paragraph(style='List Bullet')
        r_n = p.add_run(nom + ": ")
        r_n.bold = True
        r_n.font.color.rgb = COLOR_PRIMARY
        r_d = p.add_run(desc)
        r_d.font.color.rgb = COLOR_TEXT

    # =========================================================================
    # SECCIÓN 4: FLUJO TRANSACCIONAL DE MATRÍCULA
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("4. FLUJO DE DATOS Y LÓGICA DE NEGOCIO PASO A PASO")
    r.font.color.rgb = COLOR_PRIMARY

    p = doc.add_paragraph()
    p.add_run("El proceso más crítico y completo del sistema es el método ").font.color.rgb = COLOR_TEXT
    p.add_run("registrarMatricula()").bold = True
    p.add_run(" implementado en ").font.color.rgb = COLOR_TEXT
    p.add_run("MatriculaServiceImpl").bold = True
    p.add_run(". A continuación se detalla su flujo de ejecución:").font.color.rgb = COLOR_TEXT

    pasos = [
        ("Paso 1: Recepción del Payload (DTO)", "El cliente web envía un JSON con el ID del estudiante, el ID del periodo académico, un arreglo de IDs de secciones y el método de pago elegido."),
        ("Paso 2: Validación del Estudiante", "Se busca al estudiante en la base de datos y se verifica que su estado sea 'ACTIVO'. Si estuviera egresado o suspendido, se rechaza la operación con una excepción BadRequestException."),
        ("Paso 3: Validación del Periodo Académico", "Se verifica que el periodo académico exista y que su estado no sea 'CERRADO'."),
        ("Paso 4: Verificación de No Duplicidad de Matrícula", "Se consulta si el estudiante ya cuenta con una matrícula activa en el periodo seleccionado (regla de negocio: una sola matrícula por semestre)."),
        ("Paso 5: Validación de Vacantes en Tiempo Real", "Por cada sección solicitada, se verifica que (vacantes - matriculados) sea mayor a 0. Si se agotaron los cupos, se cancela la transacción."),
        ("Paso 6: Validación de No Cruce de Cursos", "Se comprueba que el alumno no esté intentando matricularse en dos secciones distintas de la misma asignatura en el mismo semestre."),
        ("Paso 7: Actualización de Cupos y Suma de Créditos", "Se incrementa el contador de alumnos matriculados en cada sección, se acumulan los créditos totales y se calcula el costo financiero total."),
        ("Paso 8: Generación de Código Único de Ficha", "Se emite un código correlativo con formato 'MAT-2026I-0001'."),
        ("Paso 9: Emisión de Pago Inmediato", "Si el usuario seleccionó un método de pago (Tarjeta, Yape, Efectivo), se inserta el registro en la tabla 'pagos' con su número de operación y la matrícula pasa a estado 'PAGADA'."),
        ("Paso 10: Persistencia Atómica (@Transactional)", "Spring Boot confirma los cambios en MySQL. Si alguna validación falla en cualquier paso, se ejecuta un Rollback automático impidiendo datos inconsistentes.")
    ]

    for title, desc in pasos:
        p = doc.add_paragraph()
        r_t = p.add_run("▶ " + title + "\n")
        r_t.bold = True
        r_t.font.color.rgb = COLOR_PRIMARY
        r_d = p.add_run(desc)
        r_d.font.color.rgb = COLOR_TEXT
        p.paragraph_format.space_after = Pt(6)

    # =========================================================================
    # SECCIÓN 5: GUION DE EXPOSICIÓN (5 A 10 MINUTOS)
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("5. GUION DE EXPOSICIÓN (ESTRUCTURA DE DEFENSA)")
    r.font.color.rgb = COLOR_PRIMARY

    p = doc.add_paragraph()
    p.add_run("Utiliza este orden cronológico para realizar una presentación fluida, profesional y contundente ante el jurado o profesor:").font.color.rgb = COLOR_TEXT

    fases_exposicion = [
        ("Minuto 0:00 - 1:30 | Introducción y Planteamiento del Problema",
         "• 'Buenos días profesor/jurado. Hoy presentaré el Sistema de Matrícula Universitaria desarrollado bajo el patrón arquitectónico MVC con Spring Boot y MySQL.'\n"
         "• 'El objetivo principal es resolver la concurrencia y asignación de cursos en periodos lectivos, gestionando vacantes en tiempo real, control de créditos y recaudación de pagos.'\n"
         "• 'Cumplimos rigurosamente con los requisitos del examen: más de 10 tablas (12 implementadas), estructura MVC estricta, control transaccional y repositorio en GitHub.'"),

        ("Minuto 1:30 - 3:30 | Arquitectura del Software (El 'Cómo está hecho')",
         "• 'Para el desarrollo elegimos una arquitectura limpia dividida en 5 capas: Modelo (Entity), Repositorio (Spring Data JPA), Servicio (Lógica de Negocio), Controlador (REST API) y Vista (Frontend SPA).'\n"
         "• 'Nuestros Controladores se encargan estrictamente de recibir y responder peticiones HTTP en formato JSON estandarizado (ApiResponse).'\n"
         "• 'Toda la lógica de negocio pesada, como el control de vacantes y la validación de doble matrícula, reside en la capa de Servicios bajo la anotación @Transactional, lo que garantiza atomicidad: o se guarda todo o no se guarda nada.'"),

        ("Minuto 3:30 - 5:30 | Modelo de Base de Datos (12 Tablas)",
         "• 'Nuestra base de datos en MySQL cuenta con 12 tablas en 3ra Forma Normal.'\n"
         "• 'Destacan carreras, planes_estudio, cursos, periodos_academicos, docentes, aulas, estudiantes, secciones, horarios, matriculas, detalles_matricula y pagos.'\n"
         "• 'Un detalle clave es la separación entre el Curso conceptual (ej. Base de Datos I) y la Sección ofertada (ej. Grupo A en turno Mañana con 30 vacantes). Esta distinción permite que una misma materia tenga múltiples profesores y horarios sin duplicar datos.'"),

        ("Minuto 5:30 - 8:30 | Demostración en Vivo del Sistema",
         "• 'Pasemos a la demostración práctica.' (Muestras el navegador en http://localhost:8080)\n"
         "• '1. Dashboard: Observamos los contadores en tiempo real (alumnos, cursos, total recaudado).'\n"
         "• '2. Registrar Alumno: Registramos un nuevo estudiante con su DNI y carrera.'\n"
         "• '3. Nueva Matrícula: Seleccionamos al alumno recién creado, elegimos el periodo 2026-I y vemos cómo el sistema carga dinámicamente solo las secciones con cupos disponibles.'\n"
         "• '4. Marcamos las asignaturas: Nótese cómo la interfaz calcula en tiempo real los créditos acumulados y el costo total.'\n"
         "• '5. Confirmamos la matrícula: El sistema genera el código único, descuenta las vacantes en MySQL y emite el comprobante de pago.'\n"
         "• '6. Mostramos en MySQL Workbench o consola que los datos se insertaron con integridad.'"),

        ("Minuto 8:30 - 10:00 | Conclusiones y Ronda de Preguntas",
         "• 'En conclusión, hemos implementado una solución robusta, mantenible y escalable, preparada para entornos de alta demanda.'\n"
         "• 'Quedo a su disposición para resolver cualquier pregunta técnica o de negocio.'")
    ]

    for titulo, contenido in fases_exposicion:
        doc.add_heading(level=2).add_run(titulo)
        p = doc.add_paragraph()
        p.add_run(contenido).font.color.rgb = COLOR_TEXT
        p.paragraph_format.space_after = Pt(8)

    # =========================================================================
    # SECCIÓN 6: BANCO DE PREGUNTAS Y RESPUESTAS DEL JURADO
    # =========================================================================
    doc.add_page_break()
    h1 = doc.add_heading(level=1)
    r = h1.add_run("6. PREGUNTAS TÍPICAS DEL PROFESOR Y RESPUESTAS RECOMENDADAS")
    r.font.color.rgb = COLOR_PRIMARY

    p = doc.add_paragraph()
    p.add_run("Estas son las preguntas más comunes que suelen hacer los docentes en una evaluación de programación web y arquitectura MVC:").font.color.rgb = COLOR_TEXT

    faq = [
        ("1. ¿Por qué separaste los Servicios en Interfaces (ICarreraService) e Implementaciones (CarreraServiceImpl)?",
         "Respuesta: Por el Principio de Inversión de Dependencias (la 'D' de SOLID). Al programar contra una interfaz y no contra una clase concreta, logramos desacoplamiento. Esto permite cambiar la implementación en el futuro (por ejemplo para pruebas unitarias con Mocks o si cambiamos de base de datos) sin tener que modificar los controladores que la consumen."),

        ("2. ¿Qué función cumple la anotación @Transactional en MatriculaServiceImpl?",
         "Respuesta: Garantiza la atomicidad de la transacción (ACID). El registro de una matrícula implica múltiples escrituras: insertar la cabecera, insertar los detalles de los cursos, restar las vacantes en cada sección y registrar el comprobante de pago. Si cualquiera de estos pasos falla, @Transactional hace un 'Rollback' automático, asegurando que la base de datos nunca quede en un estado inconsistente."),

        ("3. ¿Cómo evitaste el problema de bucle infinito (ciclos de recursión) al serializar entidades JPA a JSON?",
         "Respuesta: Lo resolvimos mediante dos mecanismos profesionales: 1) El uso de DTOs (Data Transfer Objects) que extraen únicamente los datos necesarios y planos para enviar al frontend; y 2) Las anotaciones de Jackson @JsonManagedReference en el padre y @JsonBackReference en el hijo (en la relación Matricula - DetalleMatricula), impidiendo que el serializador entre en bucle infinito."),

        ("4. ¿Cuál es la diferencia entre @Controller y @RestController en Spring?",
         "Respuesta: @Controller tradicional está pensado para vistas renderizadas en el servidor (como Thymeleaf o JSP) y retorna nombres de plantillas HTML. En cambio, @RestController combina @Controller con @ResponseBody, indicando que todos los métodos devuelven datos puros serializados directamente en JSON o XML para ser consumidos por APIs REST y clientes modernos."),

        ("5. ¿Por qué se utilizó el motor InnoDB en MySQL para las 12 tablas?",
         "Respuesta: Porque InnoDB es el motor estándar de MySQL que ofrece soporte completo para transacciones ACID, bloqueo a nivel de fila (row-level locking para evitar condiciones de carrera en alta concurrencia) y, sobre todo, integridad referencial mediante claves foráneas (Foreign Keys) con restricciones ON DELETE y ON UPDATE."),

        ("6. ¿Por qué crear una tabla 'secciones' en lugar de matricular al alumno directamente en la tabla 'cursos'?",
         "Respuesta: Por normalización de base de datos. Un curso es una definición teórica que permanece en el tiempo (ej. 'Base de Datos I'), mientras que una sección representa una oferta concreta en un semestre determinado, con un horario, un aula, un docente específico y un cupo limitado de vacantes. Si matriculáramos directamente en 'cursos', no podríamos tener múltiples turnos ni controlar vacantes por grupo."),

        ("7. ¿Qué función cumple la clase GlobalExceptionHandler con @RestControllerAdvice?",
         "Respuesta: Centraliza la captura de excepciones en toda la aplicación. En lugar de llenar los controladores con bloques try-catch repetitivos, cualquier excepción personalizada como ResourceNotFoundException o BadRequestException es interceptada por GlobalExceptionHandler, devolviendo una respuesta uniforme con formato ApiResponse y el código de estado HTTP apropiado (404 o 400)."),

        ("8. ¿Qué significa la regla 'ON UPDATE CASCADE ON DELETE RESTRICT' en las claves foráneas?",
         "Respuesta: 'ON UPDATE CASCADE' asegura que si se actualiza el identificador de un registro padre, los hijos se actualicen automáticamente. 'ON DELETE RESTRICT' impide que se borre un registro padre si existen registros dependientes (por ejemplo, impide borrar una Carrera si todavía tiene alumnos o planes de estudio asociados), protegiendo la integridad de los datos históricos."),

        ("9. ¿Cómo configuraste CORS (Cross-Origin Resource Sharing) y por qué es necesario?",
         "Respuesta: Lo configuramos mediante la clase CorsConfig implementando WebMvcConfigurer. Es indispensable en arquitecturas desacopladas porque los navegadores bloquean por seguridad las solicitudes asíncronas (fetch/ajax) enviadas desde un origen distinto al del servidor API (por ejemplo si el frontend corre en otro puerto o como archivo local)."),

        ("10. ¿Cómo se empaqueta y ejecuta este proyecto en un entorno de producción?",
         "Respuesta: Se ejecuta './mvnw.cmd package -DskipTests', lo cual compila las clases Java y genera un archivo ejecutable autocontenido (.jar) en la carpeta 'target'. Este archivo incluye un servidor web embebido (Apache Tomcat), de modo que para ponerlo en producción en cualquier servidor con Java 21 solo se ejecuta 'java -jar backend-0.0.1-SNAPSHOT.jar', sin requerir configurar un servidor externo.")
    ]

    for q, a in faq:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(8)
        p_q.paragraph_format.space_after = Pt(2)
        r_q = p_q.add_run(q)
        r_q.bold = True
        r_q.font.color.rgb = COLOR_PRIMARY
        r_q.font.size = Pt(11)

        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_before = Pt(0)
        p_a.paragraph_format.space_after = Pt(8)
        r_a = p_a.add_run(a)
        r_a.font.color.rgb = COLOR_TEXT
        r_a.font.size = Pt(10)

    # Guardar documento en el Escritorio
    desktop_path = os.path.join(os.environ["USERPROFILE"], "Desktop")
    output_filename = os.path.join(desktop_path, "GUIA_ESTUDIO_SISTEMA_MATRICULA_EXPOSICION.docx")
    doc.save(output_filename)
    print(f"Documento generado exitosamente en: {output_filename}")

if __name__ == "__main__":
    create_study_guide()
